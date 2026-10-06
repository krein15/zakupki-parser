"""Audit of a keyword set: the words file, found / junk / missed against a reference, the report and the command."""

from __future__ import annotations

from datetime import date, datetime

import pytest
from conftest import NOTICES, RegionSite
from openpyxl import load_workbook

from zkparser import __main__ as cli
from zkparser.audit import Keywords, parse_keywords, run_audit
from zkparser.cache import NoticeCache
from zkparser.excel import export_audit, headline
from zkparser.profiles import ProfileError, from_template, load_templates

TYUMEN = "72000000000"
DAY = date(2026, 10, 5)
FUEL, STORAGE = "0167100004126000028", "0167100002326000043"
REFERENCE = from_template(next(t for t in load_templates() if t.name == "Топливо и ГСМ"), (TYUMEN,))


def test_words_file():
    text = """# слова клиента
бензин, дизельное топливо
-аренда; - ремонт*
бензин
"горюче-смазочные материалы"
"""
    assert parse_keywords(text) == Keywords(("бензин", "дизельное топливо", '"горюче-смазочные материалы"'),
                                            ("аренда", "ремонт*"))


def test_words_file_without_words_is_an_error():
    with pytest.raises(ProfileError, match=r"слова\.txt: нет ни одного ключевого слова"):
        parse_keywords("# пусто\n-ремонт\n", "слова.txt")


def audit(tmp_path, words: str):
    with NoticeCache(tmp_path / "cache") as cache:
        return run_audit(RegionSite({TYUMEN: list(NOTICES)}), cache, parse_keywords(words), REFERENCE, DAY, DAY)


def test_found_junk_and_missed(tmp_path):
    result = audit(tmp_path, "бензин, хранение")
    assert [row.found.notice.reg_number for row in result.both] == [FUEL]
    assert [row.found.notice.reg_number for row in result.junk] == [STORAGE]
    assert result.missed == []
    assert (result.checked, result.found, result.relevant) == (5, 2, 1)
    assert (result.precision, result.recall) == (0.5, 1.0)
    stats = {stat.word: (stat.found, stat.relevant, stat.junk) for stat in result.word_stats()}
    assert stats == {"хранение": (1, 0, 1), "бензин": (1, 1, 0)}
    assert result.word_stats()[0].word == "хранение"  # the noisiest first


def test_missed_notices_tell_what_the_set_lacks(tmp_path):
    result = audit(tmp_path, "хранение")
    assert [row.found.notice.reg_number for row in result.missed] == [FUEL]
    assert result.missed_sum() == 2359233
    labels = [label for label, count, _ in result.hints()]
    assert "«топливо»" in labels and "ОКПД2 19.20.21.300" in labels
    assert headline(result) == (
        "Из 1 закупки, найденной вашими словами, не подходит ни одна. "
        "Ещё 1 подходящую закупку на 2 359 233,00 ₽ ваш набор пропустил."
    )


def test_the_client_set_keeps_the_reference_filters(tmp_path):
    with NoticeCache(tmp_path / "cache") as cache:
        expensive = REFERENCE.__class__(**{**REFERENCE.__dict__, "price_from": 2_000_000})
        result = run_audit(RegionSite({TYUMEN: list(NOTICES)}), cache, parse_keywords("бензин, хранение"),
                           expensive, DAY, DAY)
    assert [row.found.notice.reg_number for row in result.both] == [FUEL]
    assert result.junk == []  # storage at 1 mln ₽ is below the price of the audit for both sets


def test_report(tmp_path):
    result = audit(tmp_path, "бензин, хранение, томатн*")
    path = export_audit(tmp_path / "audit.xlsx", result, datetime(2026, 10, 6, 12, 0))
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Итоги", "Пропущено", "Мусор", "Совпало", "Слова"]
    summary = workbook["Итоги"]
    assert summary["A4"].value == (
        "Из 3 закупок, найденных вашими словами, подходят 1. Пропущенных подходящих закупок нет."
    )
    pairs = {row[0].value: row[1].value for row in summary.iter_rows() if row[0].value and row[1].value}
    assert pairs["Нашёл ваш набор"] == 3
    assert pairs["— из них подходят"] == "1 (точность 33%)"
    assert pairs["хранение"] == "1 из 1 — мусор: уберите или уточните"
    junk = workbook["Мусор"]
    assert junk["A2"].hyperlink.target.endswith("regNumber=" + STORAGE)  # the biggest money first
    assert [c.value for c in junk[1]][-1] == "Какое ваше слово сработало"
    assert workbook["Пропущено"]["A2"].value == "Пропущенных нет: ваш набор нашёл всё, что нашёл эталон"
    words = [[c.value for c in row] for row in workbook["Слова"].iter_rows(min_row=2)]
    assert [word[:4] for word in words] == [["томатн*", 1, 0, 1], ["хранение", 1, 0, 1], ["бензин", 1, 1, 0]]


class StatsSite(RegionSite):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.stats = type("Stats", (), {"requests": 0, "throttled": 0})()


def test_audit_command(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(cli, "SiteClient", lambda: StatsSite({TYUMEN: list(NOTICES)}))
    words = tmp_path / "слова.txt"
    words.write_text("хранение\n", encoding="utf-8")
    code = cli.main(["audit", str(words), "--template", "топливо", "-r", "72", "--cache-dir", str(tmp_path / "c"),
                     "--out", str(tmp_path)])
    out = capsys.readouterr().out
    assert code == 0
    assert "Аудит: 1 слов и 0 минус-слов против «Топливо и ГСМ»" in out
    assert "ваш набор пропустил" in out
    assert len(list(tmp_path.glob("Аудит слов — Топливо и ГСМ_*.xlsx"))) == 1


def test_audit_command_needs_regions_for_a_template(tmp_path):
    words = tmp_path / "слова.txt"
    words.write_text("бензин\n", encoding="utf-8")
    with pytest.raises(SystemExit):
        cli.main(["audit", str(words), "--template", "топливо"])
