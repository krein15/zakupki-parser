"""Export of open notices: the search window, the deadline check, details from XML, Excel and the command."""

from __future__ import annotations

from datetime import date, datetime, timedelta, timezone

import pytest
from conftest import FIXTURES, NOTICES
from openpyxl import load_workbook

from zkparser import __main__ as cli
from zkparser.cache import NoticeCache
from zkparser.display import days_left, price_range, regions_count, short_money
from zkparser.excel import DETAILED_COLUMNS, OPEN_COLUMNS, export_file_name, export_open_notices
from zkparser.export import LOOKBACK_DAYS, ExportQuery, export_open
from zkparser.fetch import STOPPED
from zkparser.website.notice import NOTICE_XML_PATH
from zkparser.website.search import SEARCH_PATH

TYUMEN, KHMAO = "72000000000", "86000000000"
MSK = timezone(timedelta(hours=3))
NOW = datetime(2026, 10, 8, 12, 0, tzinfo=MSK)  # 14:00 in Tyumen: applications until 08.10 08:00 are over
GENERATED = datetime(2026, 10, 8, 12, 0)

QUOTATION, DRUGS, JOINT, FUEL, AUDIT = (
    "0167100002326000043", "0167200003426008040", "0167200003426008053", "0167100004126000028", "1200700110126000001"
)
STALE = "0167200003414000001"  # still "Подача заявок" on the site since 2014


def entry(number: str, deadline: str, price: str = "1 000 000,00 ₽", title: str = "Поставка",
          organization: str = "ГКУ ТО «Центр закупок»") -> str:
    deadline_block = (f'<div class="data-block__title">Окончание подачи заявок</div>'
                      f'<div class="data-block__value">{deadline}</div>' if deadline else "")
    return (
        '<div class="search-registry-entry-block">'
        '<div class="registry-entry__header-top__title">44-ФЗ Электронный аукцион</div>'
        f'<div class="registry-entry__header-mid__number"><a href="/epz/order/notice/ea20/view/common-info.html'
        f'?regNumber={number}">№ {number}</a></div>'
        '<div class="registry-entry__header-mid__title">Подача заявок</div>'
        '<div class="registry-entry__body-block"><div class="registry-entry__body-title">Объект закупки</div>'
        f'<div class="registry-entry__body-value">{title}</div></div>'
        '<div class="registry-entry__body-block"><div class="registry-entry__body-title">Заказчик</div>'
        f'<div class="registry-entry__body-href"><a href="/x?organizationCode=01672000034">{organization}</a></div>'
        f'</div><div class="price-block__value">{price}</div>'
        '<div class="data-block__title">Размещено</div><div class="data-block__value">01.10.2026</div>'
        '<div class="data-block__title">Обновлено</div><div class="data-block__value">02.10.2026</div>'
        f"{deadline_block}</div>"
    )


class OpenSite:
    """One page of search results per region; notice XML from the fixtures."""

    def __init__(self, by_region: dict[str, list[str]]):
        self.by_region = by_region
        self.calls: list[tuple[str, dict]] = []

    def get(self, path, params=None):
        self.calls.append((path, params))
        if path == NOTICE_XML_PATH:
            return (FIXTURES / f"notice_{NOTICES[params['regNumber']]}.xml").read_bytes()
        assert path == SEARCH_PATH
        blocks = self.by_region.get(params["customerPlace"], [])
        return (f'<html><body><div class="search-results__total">{len(blocks)} записей</div>{"".join(blocks)}'
                "</body></html>").encode()

    def searches(self) -> list[dict]:
        return [params for path, params in self.calls if path == SEARCH_PATH]

    def downloads(self) -> list[str]:
        return [params["regNumber"] for path, params in self.calls if path == NOTICE_XML_PATH]


TYUMEN_RESULTS = [
    entry(STALE, "20.01.2014", "2 000 000,00 ₽"),
    entry(QUOTATION, "07.10.2026", "1 000 000,00 ₽"),  # closed yesterday
    entry(DRUGS, "08.10.2026", "233 984,40 ₽"),  # last day today
    entry(JOINT, "08.10.2026", "390 592,00 ₽"),
    entry(FUEL, "09.10.2026", "2 359 233,00 ₽"),
    entry(AUDIT, "", "1 359 672,00 ₽"),  # no deadline in the results: kept, not dropped
]


def run(site, tmp_path, query=None, **options):
    query = query or ExportQuery((TYUMEN,), 100000, 5000000)
    with NoticeCache(tmp_path / "cache") as cache:
        return export_open(site, cache, query, NOW, **options)


def test_search_asks_for_open_stage_price_and_recent_notices(tmp_path):
    site = OpenSite({})
    run(site, tmp_path, ExportQuery((TYUMEN, KHMAO), 100000, 5000000, text="бумага"))
    first, second = site.searches()
    assert (first["customerPlace"], second["customerPlace"]) == (TYUMEN, KHMAO)
    assert first["af"] == "on" and "ca" not in first
    assert (first["priceFromGeneral"], first["priceToGeneral"]) == ("100000", "5000000")
    assert first["searchString"] == "бумага"
    assert first["publishDateFrom"] == f"{NOW.date() - timedelta(days=LOOKBACK_DAYS):%d.%m.%Y}"
    assert first["publishDateTo"] == "08.10.2026"


def test_only_notices_still_taking_applications(tmp_path):
    result = run(OpenSite({TYUMEN: TYUMEN_RESULTS}), tmp_path)
    numbers = [item.hit.reg_number for item in result.notices]
    # The nearest deadline first, on the same day the most expensive on top; no deadline last.
    assert numbers == [JOINT, DRUGS, FUEL, AUDIT]
    assert (result.listed, result.closed) == (6, 2)
    assert (result.regions[TYUMEN].listed, result.regions[TYUMEN].open) == (6, 4)
    assert result.notices[0].days_left(result.today) == 0
    assert result.notices[0].notice is None  # no details asked: no downloads
    assert not result.error


def test_a_joint_purchase_in_two_regions_is_listed_once(tmp_path):
    site = OpenSite({TYUMEN: [entry(FUEL, "09.10.2026")], KHMAO: [entry(FUEL, "09.10.2026")]})
    result = run(site, tmp_path, ExportQuery((TYUMEN, KHMAO)))
    assert [(item.hit.reg_number, item.regions) for item in result.notices] == [(FUEL, (TYUMEN, KHMAO))]
    assert result.listed == 1


def test_details_download_only_open_notices_and_use_the_exact_deadline(tmp_path):
    site = OpenSite({TYUMEN: TYUMEN_RESULTS})
    result = run(site, tmp_path, ExportQuery((TYUMEN,), details=True))
    assert sorted(site.downloads()) == sorted([DRUGS, JOINT, FUEL, AUDIT])
    # Applications for DRUGS and JOINT closed at 08:00 Tyumen time today: the XML says so, the results do not.
    assert [item.hit.reg_number for item in result.notices] == [FUEL, AUDIT]
    assert result.closed == 4
    fuel = result.notices[0]
    assert fuel.notice.customers[0].customer.inn == "7202215241"
    assert fuel.deadline_day == date(2026, 10, 9)


def test_stop_keeps_what_was_collected(tmp_path):
    site = OpenSite({TYUMEN: TYUMEN_RESULTS})
    result = run(site, tmp_path, ExportQuery((TYUMEN, KHMAO)), cancel=lambda: len(site.searches()) > 1)
    assert result.error == STOPPED
    assert [item.hit.reg_number for item in result.notices] == [JOINT, DRUGS, FUEL, AUDIT]
    assert list(result.regions) == [TYUMEN]


def test_listing_reports_progress(tmp_path):
    seen = []
    run(OpenSite({TYUMEN: TYUMEN_RESULTS}), tmp_path, listing=lambda *args: seen.append(args))
    assert seen == [(TYUMEN, 6, 6)]


def test_excel_from_search_results(tmp_path):
    result = run(OpenSite({TYUMEN: TYUMEN_RESULTS}), tmp_path)
    path = export_open_notices(tmp_path / "open.xlsx", result, GENERATED)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Закупки", "Выгрузка"]
    sheet = workbook["Закупки"]
    assert [c.value for c in sheet[1]] == [column.title for column in OPEN_COLUMNS]
    row = [c.value for c in sheet[2]]
    assert row[:5] == [JOINT, "Поставка", 390592, datetime(2026, 10, 8), 0]
    assert row[6] == "ГКУ ТО «Центр закупок»"
    assert sheet["A2"].hyperlink.target.endswith(JOINT)
    summary = {r[0].value: r[1].value for r in workbook["Выгрузка"].iter_rows() if r[0].value}
    assert summary["На сайте с этапом «Подача заявок»"] == 6
    assert summary["Из них срок подачи уже прошёл"] == 2
    assert summary["Открыты — в выгрузке"] == 4
    assert summary["Начальная цена"] == "от 100 000,00 ₽ до 5 000 000,00 ₽"


def test_excel_with_details_has_customers_and_positions(tmp_path):
    result = run(OpenSite({TYUMEN: TYUMEN_RESULTS}), tmp_path, ExportQuery((TYUMEN,), details=True))
    workbook = load_workbook(export_open_notices(tmp_path / "open.xlsx", result, GENERATED))
    assert workbook.sheetnames == ["Закупки", "Позиции", "Выгрузка"]
    sheet = workbook["Закупки"]
    assert [c.value for c in sheet[1]] == [column.title for column in DETAILED_COLUMNS]
    row = {column.title: cell.value for column, cell in zip(DETAILED_COLUMNS, sheet[2], strict=True)}
    assert row["Подача заявок до"] == datetime(2026, 10, 9, 10, 0)
    assert row["Часовой пояс"] == "МСК+2"
    assert row["ИНН заказчика"] == "7202215241"
    positions = workbook["Позиции"]
    assert "Подошла" not in [c.value for c in positions[1]]
    assert positions.max_row == 1 + 4 + 1  # header, four fuel positions, one of the audit


def test_empty_export_says_so(tmp_path):
    result = run(OpenSite({}), tmp_path)
    sheet = load_workbook(export_open_notices(tmp_path / "open.xlsx", result, GENERATED))["Закупки"]
    assert sheet["A2"].value == "Открытых закупок нет"


def test_export_file_name(tmp_path):
    result = run(OpenSite({}), tmp_path)
    assert export_file_name(result, GENERATED) == (
        "Открытые закупки — Тюменская область, 100 тыс. – 5 млн ₽_2026-10-08_12-00-00.xlsx"
    )
    several = run(OpenSite({}), tmp_path, ExportQuery((TYUMEN, KHMAO)))
    assert export_file_name(several, GENERATED) == "Открытые закупки — 2 региона_2026-10-08_12-00-00.xlsx"


@pytest.mark.parametrize(
    ("value", "text"),
    [(950, "950"), (300_000, "300 тыс."), (1_000_000, "1 млн"), (1_500_000, "1,5 млн"), (2_000_000_000, "2 млрд")],
)
def test_short_money(value, text):
    assert short_money(value) == text


@pytest.mark.parametrize(
    ("count", "text"), [(1, "1 регион"), (2, "2 региона"), (5, "5 регионов"), (11, "11 регионов"), (21, "21 регион"),
                        (22, "22 региона"), (14, "14 регионов")],
)
def test_regions_count(count, text):
    assert regions_count(count) == text


@pytest.mark.parametrize(
    ("days", "text"),
    [(0, "последний день"), (1, "остался 1 день"), (3, "осталось 3 дня"), (5, "осталось 5 дней"),
     (21, "остался 21 день"), (11, "осталось 11 дней")],
)
def test_days_left(days, text):
    assert days_left(days) == text


def test_price_range():
    assert price_range(None, None) == ""
    assert price_range(1000, None) == "от 1 000,00 ₽"
    assert price_range(None, 5_000_000, short=True) == "до 5 млн ₽"
    assert price_range(1_000_000, 5_000_000, short=True) == "1 млн – 5 млн ₽"


def test_export_command(tmp_path, monkeypatch, capsys):
    site = OpenSite({TYUMEN: TYUMEN_RESULTS})
    monkeypatch.setattr(cli, "SiteClient", lambda: StatsSite(site))
    code = cli.main(["export", "-r", "72", "--price-from", "100000", "--price-to", "5000000",
                     "--cache-dir", str(tmp_path / "cache"), "--out", str(tmp_path)])
    out = capsys.readouterr().out
    assert code == 0
    assert "Тюменская область: с этапом «Подача заявок» 6, из них открыто" in out
    assert FUEL in out and STALE not in out
    assert len(list(tmp_path.glob("Открытые закупки — Тюменская область*.xlsx"))) == 1


def test_export_command_checks_the_price_range():
    with pytest.raises(SystemExit):
        cli.main(["export", "-r", "72", "--price-from", "5000000", "--price-to", "1000000"])


class StatsSite:
    """The fake site with the request counters of SiteClient, which the command prints."""

    def __init__(self, site):
        self.site = site
        self.stats = type("Stats", (), {"requests": 0, "throttled": 0})()

    def get(self, path, params=None):
        return self.site.get(path, params)
