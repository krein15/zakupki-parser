"""Excel report: sheets, values and formats, links, highlighted positions, safety of notice text."""

from __future__ import annotations

from dataclasses import replace
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pytest
from conftest import FIXTURES
from openpyxl import load_workbook

from zkparser.__main__ import write_reports
from zkparser.excel import NOTICE_COLUMNS, POSITION_COLUMNS, build_file_name, export_profile, moscow_offset
from zkparser.matching import Verdict
from zkparser.notice_xml import parse_notice
from zkparser.pipeline import Found, ProfileResult
from zkparser.profiles import Profile
from zkparser.website.search import SearchHit

TYUMEN = "72000000000"
DAY = date(2026, 9, 30)
GENERATED = datetime(2026, 10, 3, 12, 30)
PROFILE = Profile(
    name='Продукты/клиент: "А"', regions=(TYUMEN,), keywords=("томатн*",), price_from=10000, path=Path("a.toml")
)


def found(
    name: str, reasons=("название: «томатн*»",), positions=(1,), stage="Подача заявок", **changes
) -> Found:
    notice = replace(parse_notice((FIXTURES / f"notice_{name}.xml").read_bytes()), **changes)
    hit = SearchHit(reg_number=notice.reg_number, url=notice.url, published=date(2026, 10, 2), stage=stage)
    return Found(notice, hit, Verdict(True, reasons, positions), (TYUMEN,))


def export(tmp_path, matches, checked=10):
    path = tmp_path / "report.xlsx"
    export_profile(path, ProfileResult(PROFILE, checked, matches), DAY, DAY, GENERATED)
    return load_workbook(path)


def test_sheets_and_headers(tmp_path):
    workbook = export(tmp_path, [found("joint")])
    assert workbook.sheetnames == ["Закупки", "Позиции", "Профиль"]
    assert [c.value for c in workbook["Закупки"][1]] == [column.title for column in NOTICE_COLUMNS]
    assert [c.value for c in workbook["Позиции"][1]] == [column.title for column in POSITION_COLUMNS]
    assert workbook["Закупки"].freeze_panes == "B2"


def test_notice_row(tmp_path):
    reasons = ("название: «томатн*»", "позиция 1 «Томатная паста»: «томатн*»")
    row = {cell.column_letter: cell for cell in export(tmp_path, [found("joint", reasons)])["Закупки"][2]}
    assert row["A"].value == "0167200003426008053"
    assert row["A"].hyperlink.target.endswith("regNumber=0167200003426008053")
    assert row["B"].value == "поставка продуктов питания (томатная паста)"
    assert row["C"].value == "название: «томатн*»\nпозиция 1 «Томатная паста»: «томатн*»"
    assert row["D"].value == 390592
    assert row["D"].number_format == '#,##0.00 "₽"'
    assert row["E"].value == datetime(2026, 10, 8, 8, 0)  # local time of the customer, Excel keeps no zones
    assert row["F"].value == "МСК+2"
    assert row["G"].value == "Подача заявок"
    assert row["H"].value == "Электронный аукцион"
    assert len(row["I"].value.splitlines()) == 4
    assert row["J"].value.splitlines() == ["7224009250", "7216001666", "7228000177", "7215004008"]
    assert row["K"].value == "Тюменская область"
    assert row["L"].value == "АО «Сбербанк-АСТ»"
    assert row["M"].value == datetime(2026, 10, 2)  # from the search, not the XML's signing date
    assert row["N"].value == "совместная закупка, заказчиков: 4"


@pytest.mark.parametrize(
    ("stage", "color"),
    [
        ("Подача заявок", "15803D"),
        ("Работа комиссии", "B45309"),
        ("Определение поставщика отменено", "B91C1C"),
        ("Закупка завершена", "6B7280"),
    ],
)
def test_stage_is_coloured(tmp_path, stage, color):
    cell = export(tmp_path, [found("drugs", stage=stage)])["Закупки"]["G2"]
    assert cell.value == stage
    assert cell.font.color.rgb.endswith(color)


def test_positions_sheet_marks_matching_positions(tmp_path):
    sheet = export(tmp_path, [found("auction_ktru", positions=(2, 3))])["Позиции"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 4
    assert [row[1] for row in rows] == [1, 2, 3, 4]
    assert [row[3] for row in rows] == [None, "да", "да", None]
    assert sheet["C2"].value == "Топливо дизельное (розничная реализация)"
    assert (sheet["E2"].value, sheet["F2"].value) == ("19.20.21.300", "19.20.21.300-00000009")
    assert (sheet["G2"].value, sheet["H2"].value, sheet["I2"].value, sheet["J2"].value) == (
        800, "Литр; кубический дециметр", 91.58, 73264
    )
    assert sheet["A3"].fill.fgColor.rgb.endswith("E8F5E9")
    assert not sheet["A2"].fill.fgColor.rgb.endswith("E8F5E9")


def test_undefined_quantity_keeps_unit_prices_only(tmp_path):
    workbook = export(tmp_path, [found("quotation")])
    assert "количество не определено" in workbook["Закупки"]["N2"].value
    quantity, _, price, total = next(workbook["Позиции"].iter_rows(min_row=2, min_col=7, values_only=True))
    assert (quantity, price, total) == (None, 25.94, None)


def test_notice_text_never_becomes_a_formula(tmp_path):
    evil = found("drugs", title='=HYPERLINK("http://example.org","клик")\x07')
    cell = export(tmp_path, [evil])["Закупки"]["B2"]
    assert cell.data_type == "s"
    assert cell.value == '=HYPERLINK("http://example.org","клик")'


def test_empty_report_says_so(tmp_path):
    workbook = export(tmp_path, [], checked=59)
    assert workbook["Закупки"]["A2"].value == "Подходящих закупок нет"
    values = list(workbook["Профиль"].iter_rows(values_only=True))
    assert ("Проверено закупок", 59) in values
    assert ("Подошло", 0) in values


def test_profile_sheet(tmp_path):
    values = dict(row[:2] for row in export(tmp_path, [found("joint")])["Профиль"].iter_rows(values_only=True))
    assert values["Регионы"] == "Тюменская область"
    assert values["Период размещения"] == "30.09.2026"
    assert values["Ключевые слова"] == "томатн*"
    assert values["Начальная цена"] == "от 10 000,00 ₽"
    assert values["Этап"] == "подача заявок"


@pytest.mark.parametrize(("hours", "text"), [(5, "МСК+2"), (3, "МСК"), (2, "МСК-1"), (12, "МСК+9")])
def test_moscow_offset(hours, text):
    assert moscow_offset(datetime(2026, 10, 1, tzinfo=timezone(timedelta(hours=hours)))) == text


def test_moscow_offset_without_time():
    assert moscow_offset(None) == ""


def test_file_name_is_safe_for_windows():
    name = build_file_name(PROFILE, GENERATED)
    assert name == "Продукты клиент А_2026-10-03_12-30-00.xlsx"


def test_profiles_with_the_same_name_get_separate_files(tmp_path):
    results = [ProfileResult(PROFILE), ProfileResult(PROFILE)]
    first, second = write_reports(results, DAY, DAY, tmp_path)
    assert first != second
    assert second.stem.endswith("(2)")
    assert first.exists() and second.exists()
