"""Excel reports: one file per profile, so a client gets a report of their own, and the export of open notices.

* "Закупки" — a row per matching notice: number (a link to the notice), title, why it matched, price, application
  deadline, current stage, customers and their INNs, region, platform, notes.
* "Позиции" — every position of those notices, the matching ones highlighted.
* "Профиль" — what was searched for, over which period and how many notices were checked.

The export of open notices (export.py) has the same "Закупки" and "Позиции" sheets without the matching, and
"Выгрузка" with what was exported and how many notices the site still listed after their deadline.

Times are the customer's local time, as the notice gives them; "Часовой пояс" shows the offset from Moscow the way
the EIS site does (МСК+2). Text from notices never becomes a formula.
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .display import price_range, regions_count
from .export import ExportResult, OpenNotice
from .models import Notice
from .pipeline import Found, ProfileResult
from .profiles import Profile
from .regions import REGIONS

HEADER_FILL = PatternFill("solid", fgColor="2B2D42")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="2B2D42")
SECTION_FONT = Font(bold=True, size=12, color="2B2D42")
MUTED_FONT = Font(color="6B7280")
LINK_FONT = Font(color="0563C1", underline="single")
MATCH_FILL = PatternFill("solid", fgColor="E8F5E9")
# The stage at report time, coloured by a word it contains.
STAGE_FONTS = {
    "подача": Font(bold=True, color="15803D"),  # Подача заявок
    "окончен": Font(bold=True, color="B45309"),  # Приём заявок окончен (the site still says "Подача заявок")
    "комисси": Font(bold=True, color="B45309"),  # Работа комиссии
    "отмен": Font(bold=True, color="B91C1C"),  # Определение поставщика отменено, Закупка отменена
    "заверш": Font(color="6B7280"),  # Определение поставщика завершено, Закупка завершена
}
STAGE_COLUMN = "Этап"
TABLE_STYLE = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)

MONEY = '#,##0.00 "₽"'
DATETIME = "DD.MM.YYYY HH:MM"
DATE = "DD.MM.YYYY"
QUANTITY = "#,##0.###"
MAX_CELL_TEXT = 32_000
MAX_ROW_HEIGHT = 150
LINE_HEIGHT = 15
MOSCOW_UTC_OFFSET = 3


@dataclass(frozen=True)
class Column:
    title: str
    width: int
    number_format: str | None = None
    wrap: bool = False


NOTICE_COLUMNS = [
    Column("№ закупки", 22),
    Column("Название", 48, wrap=True),
    Column("Почему подошла", 48, wrap=True),
    Column("Начальная цена", 17, MONEY),
    Column("Подача заявок до", 17, DATETIME),
    Column("Часовой пояс", 10),
    Column("Этап", 18, wrap=True),
    Column("Способ", 24, wrap=True),
    Column("Заказчик", 40, wrap=True),
    Column("ИНН заказчика", 14, wrap=True),
    Column("Регион", 22, wrap=True),
    Column("Площадка", 22, wrap=True),
    Column("Размещено", 12, DATE),
    Column("Примечание", 32, wrap=True),
]
POSITION_COLUMNS = [
    Column("№ закупки", 22),
    Column("№", 5),
    Column("Наименование", 50, wrap=True),
    Column("Подошла", 9),
    Column("ОКПД2", 14),
    Column("КТРУ", 24),
    Column("Количество", 12, QUANTITY),
    Column("Ед. изм.", 14, wrap=True),
    Column("Цена за ед.", 15, MONEY),
    Column("Сумма", 16, MONEY),
]
MATCHED_POSITION = "Подошла"
# The export of open notices: from the search results only, or with the details of each notice's XML.
OPEN_COLUMNS = [
    Column("№ закупки", 22),
    Column("Название", 56, wrap=True),
    Column("Начальная цена", 17, MONEY),
    Column("Подача заявок до", 13, DATE),
    Column("Осталось дней", 10),
    Column("Способ", 24, wrap=True),
    Column("Заказчик или организатор", 44, wrap=True),
    Column("Регион", 22, wrap=True),
    Column("Размещено", 12, DATE),
    Column("Обновлено", 12, DATE),
]
DETAILED_COLUMNS = [
    Column("№ закупки", 22),
    Column("Название", 56, wrap=True),
    Column("Начальная цена", 17, MONEY),
    Column("Подача заявок до", 17, DATETIME),
    Column("Часовой пояс", 10),
    Column("Осталось дней", 10),
    Column("Способ", 24, wrap=True),
    Column("Заказчик", 40, wrap=True),
    Column("ИНН заказчика", 14, wrap=True),
    Column("Регион", 22, wrap=True),
    Column("Площадка", 22, wrap=True),
    Column("Размещено", 12, DATE),
    Column("Примечание", 32, wrap=True),
]


def build_file_name(profile: Profile, generated_at: datetime) -> str:
    return f"{_file_subject(profile)}_{generated_at:%Y-%m-%d_%H-%M-%S}.xlsx"


def daily_file_name(profile: Profile, day: date) -> str:
    """Monitoring rewrites one report a day instead of piling up a file an hour."""
    return f"{_file_subject(profile)}_мониторинг_{day:%Y-%m-%d}.xlsx"


def export_file_name(result: ExportResult, generated_at: datetime) -> str:
    """"Открытые закупки — Тюменская область, 1 млн – 5 млн ₽_2026-10-05_14-30-00.xlsx"."""
    regions = result.query.regions
    subject = REGIONS.get(regions[0], regions[0]) if len(regions) == 1 else regions_count(len(regions))
    prices = price_range(result.query.price_from, result.query.price_to, short=True)
    name = f"Открытые закупки — {subject}" + (f", {prices}" if prices else "")
    return f"{_safe(name)}_{generated_at:%Y-%m-%d_%H-%M-%S}.xlsx"


def _file_subject(profile: Profile) -> str:
    return _safe(profile.name)[:60].strip() or "закупки"


def _safe(name: str) -> str:
    return " ".join(re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", name).split())


def export_profile(path: Path, result: ProfileResult, start: date, end: date, generated_at: datetime) -> Path:
    workbook = Workbook()
    notices = workbook.active
    notices.title = "Закупки"
    _write_table(
        notices,
        "Notices",
        NOTICE_COLUMNS,
        [_notice_row(found, generated_at.astimezone()) for found in result.matches],
        links=[found.notice.url for found in result.matches],
        empty="Подходящих закупок нет",
    )

    _write_positions(workbook.create_sheet("Позиции"),
                     [(found.notice, set(found.verdict.positions)) for found in result.matches], matching=True)
    _write_profile(workbook.create_sheet("Профиль"), result, start, end, generated_at)
    workbook.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def export_open_notices(path: Path, result: ExportResult, generated_at: datetime) -> Path:
    workbook = Workbook()
    notices = workbook.active
    notices.title = "Закупки"
    detailed = result.query.details
    row = _detailed_row if detailed else _open_row
    _write_table(
        notices,
        "Notices",
        DETAILED_COLUMNS if detailed else OPEN_COLUMNS,
        [row(item, result.today) for item in result.notices],
        links=[item.hit.url for item in result.notices],
        empty="Открытых закупок нет",
    )
    if detailed:
        _write_positions(workbook.create_sheet("Позиции"),
                         [(item.notice, set()) for item in result.notices if item.notice], matching=False)
    _write_export(workbook.create_sheet("Выгрузка"), result, generated_at)
    workbook.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _write_positions(sheet: Worksheet, notices: list[tuple[Notice, set[int]]], *, matching: bool) -> None:
    """Every position of the notices; for a profile a column marks and a fill highlights the matching ones."""
    columns = POSITION_COLUMNS if matching else [c for c in POSITION_COLUMNS if c.title != MATCHED_POSITION]
    rows, links, highlighted = [], [], []
    for notice, matched in notices:
        for number, position in enumerate(notice.positions, 1):
            undefined = notice.quantity_undefined
            hit = number in matched
            rows.append(
                [
                    notice.reg_number,
                    number,
                    position.name,
                    *(["да" if hit else ""] if matching else []),
                    position.okpd2_code,
                    position.ktru_code,
                    None if undefined else position.quantity,
                    position.unit,
                    position.price,
                    None if undefined else position.total,  # with no quantity set it repeats the unit price
                ]
            )
            links.append(notice.url)
            highlighted.append(hit)
    _write_table(sheet, "Positions", columns, rows, links=links, highlighted=highlighted if matching else None)


def moscow_offset(moment: datetime | None) -> str:
    """+05:00 → "МСК+2", +03:00 → "МСК", +02:00 → "МСК-1"."""
    offset = moment.utcoffset() if moment else None
    if offset is None:
        return ""
    hours = offset.total_seconds() / 3600 - MOSCOW_UTC_OFFSET
    return "МСК" if hours == 0 else f"МСК{hours:+g}"


def stage_font(stage: str) -> Font:
    lowered = stage.casefold()
    return next((font for word, font in STAGE_FONTS.items() if word in lowered), Font())


def _notice_row(found: Found, now: datetime) -> list[Any]:
    notice = found.notice
    customers = [part.customer for part in notice.customers]
    return [
        notice.reg_number,
        notice.title,
        "\n".join(found.verdict.reasons),
        notice.max_price,
        _local(notice.applications_end),
        moscow_offset(notice.applications_end),
        found.stage(now),
        notice.placing_way,
        "\n".join(customer.name for customer in customers),
        "\n".join(customer.inn for customer in customers if customer.inn),
        ", ".join(REGIONS.get(code, code) for code in found.regions),
        notice.etp,
        # The search shows when the notice actually went public; the XML's publishDTInEIS is when it was signed,
        # which is earlier for a notice scheduled to appear later (plannedPublishDate).
        found.hit.published or (notice.published.date() if notice.published else None),
        _notes(notice),
    ]


def _notes(notice: Notice) -> str:
    notes = []
    if len(notice.customers) > 1:
        notes.append(f"совместная закупка, заказчиков: {len(notice.customers)}")
    if notice.quantity_undefined:
        notes.append("количество не определено: начальная цена — предел контракта, в позициях цены за единицу")
    if notice.version > 1:
        notes.append(f"редакция {notice.version}")
    if notice.currency not in ("", "RUB"):
        notes.append(f"валюта {notice.currency}")
    return "\n".join(notes)


def _open_row(item: OpenNotice, today: date) -> list[Any]:
    hit = item.hit
    return [
        hit.reg_number,
        hit.title,
        hit.price,
        hit.deadline,
        item.days_left(today),
        hit.placing_way,
        hit.organization,
        _regions(item.regions),
        hit.published,
        hit.updated,
    ]


def _detailed_row(item: OpenNotice, today: date) -> list[Any]:
    hit, notice = item.hit, item.notice
    if notice is None:  # the XML did not download: what the search showed
        return [hit.reg_number, hit.title, hit.price, hit.deadline, "", item.days_left(today), hit.placing_way,
                hit.organization, "", _regions(item.regions), "", hit.published, "подробности не скачались"]
    customers = [part.customer for part in notice.customers]
    return [
        notice.reg_number,
        notice.title,
        notice.max_price,
        _local(notice.applications_end) or hit.deadline,
        moscow_offset(notice.applications_end),
        item.days_left(today),
        notice.placing_way,
        "\n".join(customer.name for customer in customers),
        "\n".join(customer.inn for customer in customers if customer.inn),
        _regions(item.regions),
        notice.etp,
        hit.published or (notice.published.date() if notice.published else None),
        _notes(notice),
    ]


def _regions(codes: tuple[str, ...]) -> str:
    return ", ".join(REGIONS.get(code, code) for code in codes)


def _local(moment: datetime | None) -> datetime | None:
    """Excel has no time zones: keep the wall-clock time of the customer's region."""
    return moment.replace(tzinfo=None) if moment else None


def _write_table(
    sheet: Worksheet,
    name: str,
    columns: list[Column],
    rows: list[list[Any]],
    *,
    links: list[str] | None = None,
    highlighted: list[bool] | None = None,
    empty: str = "",
) -> None:
    for col, column in enumerate(columns, 1):
        cell = sheet.cell(row=1, column=col, value=column.title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(col)].width = column.width
    sheet.row_dimensions[1].height = 32

    for index, values in enumerate(rows):
        row = index + 2
        lines = 1
        for col, (column, value) in enumerate(zip(columns, values, strict=True), 1):
            cell = sheet.cell(row=row, column=col)
            _set_value(cell, value)
            if column.number_format:
                cell.number_format = column.number_format
            cell.alignment = Alignment(vertical="top", wrap_text=column.wrap)
            if highlighted and highlighted[index]:
                cell.fill = MATCH_FILL
            if column.title == STAGE_COLUMN and isinstance(value, str):
                cell.font = stage_font(value)
            if column.wrap and isinstance(value, str):
                lines = max(lines, _estimate_lines(value, column.width))
        if links and links[index]:
            first = sheet.cell(row=row, column=1)
            first.hyperlink, first.font = links[index], LINK_FONT
        if lines > 1:
            sheet.row_dimensions[row].height = min(lines * LINE_HEIGHT, MAX_ROW_HEIGHT)
    if not rows and empty:
        sheet.cell(row=2, column=1, value=empty).font = MUTED_FONT

    table = Table(displayName=name, ref=f"A1:{get_column_letter(len(columns))}{max(len(rows) + 1, 2)}")
    table.tableStyleInfo = TABLE_STYLE
    sheet.add_table(table)
    sheet.freeze_panes = "B2"


def _set_value(cell: Cell, value: Any) -> None:
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)[:MAX_CELL_TEXT]
        cell.value = value or None
        if value.startswith("="):
            cell.data_type = "s"  # text from a notice must not become a formula
    else:
        cell.value = value


def _estimate_lines(text: str, width: int) -> int:
    """Excel does not fit row heights of generated files, so estimate how many lines the text wraps to."""
    chars_per_line = max(int(width * 1.15), 1)
    return sum(max(math.ceil(len(line) / chars_per_line), 1) for line in text.splitlines() or [""])


def _write_profile(sheet: Worksheet, result: ProfileResult, start: date, end: date, generated_at: datetime) -> None:
    profile = result.profile
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 90
    sheet["A1"] = f"Закупки по профилю «{profile.name}»"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = f"Сформировано {generated_at:%d.%m.%Y в %H:%M} по данным zakupki.gov.ru"
    sheet["A2"].font = MUTED_FONT

    period = f"{start:%d.%m.%Y}" if start == end else f"{start:%d.%m.%Y} – {end:%d.%m.%Y}"
    row = _section(sheet, 4, "Что искали")
    row = _pairs(
        sheet,
        row,
        [
            ("Регионы", ", ".join(REGIONS.get(code, code) for code in profile.regions)),
            ("Период размещения", period),
            ("Этап", "любой" if profile.all_stages else "подача заявок"),
            ("Ключевые слова", ", ".join(profile.keywords) or "—"),
            ("Минус-слова", ", ".join(profile.minus) or "—"),
            ("ОКПД2", ", ".join(profile.okpd2) or "—"),
            ("КТРУ", ", ".join(profile.ktru) or "—"),
            ("Только заказчики (ИНН)", ", ".join(profile.customer_inn) or "любые"),
            ("Кроме заказчиков (ИНН)", ", ".join(profile.exclude_customer_inn) or "—"),
            ("Начальная цена", price_range(profile.price_from, profile.price_to) or "любая"),
        ],
    )
    row = _section(sheet, row + 1, "Итоги")
    row = _pairs(sheet, row, [("Проверено закупок", result.checked), ("Подошло", len(result.matches))])
    row = _section(sheet, row + 1, "Как читать отчёт")
    for note in (
        "«Почему подошла» — какие слова и коды профиля нашлись: в названии или в позиции. Подошедшие позиции "
        "выделены на листе «Позиции».",
        "Срок подачи заявок — местное время заказчика; «Часовой пояс» — разница с Москвой, как на сайте ЕИС.",
        "Если количество не определено, начальная цена — предел контракта, а в позициях указаны цены за единицу.",
        "Номер закупки — ссылка на её страницу в ЕИС.",
    ):
        sheet.cell(row=row, column=1, value=note)
        row += 1


def _write_export(sheet: Worksheet, result: ExportResult, generated_at: datetime) -> None:
    query = result.query
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 90
    sheet["A1"] = "Открытые закупки 44-ФЗ"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = f"Сформировано {generated_at:%d.%m.%Y в %H:%M} по данным zakupki.gov.ru"
    sheet["A2"].font = MUTED_FONT

    row = _section(sheet, 4, "Что выгружали")
    row = _pairs(
        sheet,
        row,
        [
            ("Регионы", _regions(query.regions)),
            ("Начальная цена", price_range(query.price_from, query.price_to) or "любая"),
            ("Слова (поиск ЕИС)", query.text or "—"),
            ("Размещены", f"с {result.since:%d.%m.%Y} — за {(result.today - result.since).days} дней"),
            ("Подробности", "из извещений: заказчики, ИНН, позиции" if query.details else "нет — по выдаче поиска"),
        ],
    )
    row = _section(sheet, row + 1, "Итоги")
    pairs: list[tuple[str, Any]] = [
        ("На сайте с этапом «Подача заявок»", result.listed),
        ("Из них срок подачи уже прошёл", result.closed),
        ("Открыты — в выгрузке", len(result.notices)),
    ]
    if len(query.regions) > 1:
        pairs += [(REGIONS.get(code, code), f"открыто {count.open} из {count.listed}")
                  for code, count in result.regions.items()]
    if any(count.truncated for count in result.regions.values()):
        pairs.append(("Внимание", "сайт отдаёт не больше 5 000 закупок на запрос — часть не попала, сузьте цену"))
    if result.error:
        pairs.append(("Выгрузка остановлена", result.error))
    row = _pairs(sheet, row, pairs)

    row = _section(sheet, row + 1, "Как читать")
    notes = [
        "Сайт ЕИС не меняет этап «Подача заявок» и после окончания срока, бывает годами. Программа сама сверяет срок "
        "подачи заявок: в выгрузке только закупки, по которым заявки ещё принимаются.",
        "«Подача заявок до» — последний день по времени заказчика; «Осталось дней» — 0, если срок истекает сегодня.",
    ]
    if query.details:
        notes += [
            "«Часовой пояс» — разница с Москвой, как на сайте ЕИС.",
            "Если количество не определено, начальная цена — предел контракта, а в позициях указаны цены за единицу.",
        ]
    else:
        notes.append(
            "«Заказчик или организатор» — как в выдаче поиска: сам заказчик или уполномоченный орган, который "
            "проводит закупку за него. Настоящий заказчик, его ИНН и позиции — в выгрузке с подробностями."
        )
    if any(item.deadline_day is None for item in result.notices):
        notes.append(
            "Без срока подачи — закупки, у которых сайт его не показывает (например, малого объёма через электронный "
            "магазин, ч. 12 ст. 93): они в конце списка, открыты ли — проверьте на сайте."
        )
    notes.append("Номер закупки — ссылка на её страницу в ЕИС.")
    for note in notes:
        sheet.cell(row=row, column=1, value=note)
        row += 1


def _section(sheet: Worksheet, row: int, title: str) -> int:
    sheet.cell(row=row, column=1, value=title).font = SECTION_FONT
    return row + 1


def _pairs(sheet: Worksheet, row: int, pairs: list[tuple[str, Any]]) -> int:
    for label, value in pairs:
        sheet.cell(row=row, column=1, value=label).font = MUTED_FONT
        cell = sheet.cell(row=row, column=2)
        _set_value(cell, value)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        row += 1
    return row
